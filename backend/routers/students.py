"""Students CRUD, ledger, siblings, bulk import/delete/reassign."""
import asyncio
import re
from typing import Any, Dict, List, Optional
from datetime import datetime
from fastapi import APIRouter, HTTPException, Depends, Response
from core import (
    db, StudentIn, audit, gen_id, get_current_user, now_iso, require_roles,
    eligible_receipt_codes_for_class, compute_fee_items, apply_opening_paid,
)

router = APIRouter(prefix="/api", tags=["students"])

def _month_label(month: str) -> str:
    """'2026-09' -> 'September 2026'"""
    try:
        return datetime.strptime(month, "%Y-%m").strftime("%B %Y")
    except Exception:
        return month


@router.get("/students")
async def list_students(
    q: Optional[str] = None,
    department_id: Optional[str] = None,
    class_id: Optional[str] = None,
    class_name: Optional[str] = None,
    medium: Optional[str] = None,
    stream: Optional[str] = None,
    section: Optional[str] = None,
    bus_required: Optional[bool] = None,
    limit: int = 100,
    user = Depends(get_current_user),
):
    # 'duplicate' is a terminal status set only by the one-off duplicate-
    # student consolidation flow (a student record merged into another
    # surviving admission no.) — it must never appear in a normal listing,
    # the whole point of marking it that way instead of hard-deleting it.
    query: Dict[str, Any] = {"status": {"$ne": "duplicate"}}
    if department_id: query["department_id"] = department_id
    if class_id: query["class_id"] = class_id
    # Students screen's Medium filter (Secondary only — Class 9/10 English vs
    # Marathi). Matches the student's own stored `medium` field directly, the
    # same field every class doc's medium is kept in sync with.
    if medium: query["medium"] = medium
    # Students screen's Stream filter (Junior College only — Class 11/12
    # Arts/Commerce/Science/Bi-Focal). Matches the student's own stored
    # `stream` field directly — the same field every class doc's stream is
    # kept in sync with (verified zero mismatches app-wide). Without this,
    # `class_name` alone (below) matches every stream sharing that class
    # name, which was the actual bug: selecting "Class 11" + a stream with
    # no way to narrow further meant every 11th-grade student showed up
    # regardless of stream.
    if stream: query["stream"] = stream
    # Section filter — Classes 1-8 real A/B/C sections (and any other class
    # where section is meaningfully set). Same exact-match-on-student's-own-
    # stored-field pattern as medium/stream, so Class + Medium + Section (or
    # any combination) all AND together consistently, the same one filtering
    # system used everywhere else in this endpoint — not a class-specific
    # special case.
    if section: query["section"] = section
    # Used by the BUS receipt search so only students with a currently active
    # bus assignment are selectable — same bus_required flag already trusted
    # elsewhere in the app (e.g. the dashboard's active-bus-students count).
    if bus_required is not None: query["bus_required"] = bus_required
    if class_name:
        # Class-only filter (no section/medium/stream split): match every class doc with this
        # exact name across every department/medium/stream, e.g. "Class 9" covers both English
        # and Semi Medium 9th-grade students, and "Class 11" covers all four streams — narrow
        # further with `medium`/`stream` above when that's not wanted.
        ids = [c["id"] async for c in db.classes.find({"name": class_name}, {"_id": 0, "id": 1})]
        query["class_id"] = {"$in": ids}
    if q:
        # re.escape makes every regex metacharacter in the user's search text
        # (parentheses, brackets, +, ., *, etc.) literal, so it's matched as
        # plain text instead of being interpreted as regex syntax - otherwise
        # a search like "O (Brien" throws an uncaught MongoDB OperationFailure
        # ("missing closing parenthesis") instead of just finding no results.
        # Still a case-insensitive partial/substring match, same as before.
        safe_q = re.escape(q)
        query["$or"] = [
            {"admission_no": {"$regex": safe_q, "$options": "i"}},
            {"name": {"$regex": safe_q, "$options": "i"}},
            {"guardian_mobile": {"$regex": safe_q, "$options": "i"}},
        ]
    return await db.students.find(query, {"_id":0}).limit(limit).to_list(limit)

@router.get("/students/{sid}")
async def get_student(sid: str, user = Depends(get_current_user)):
    s = await db.students.find_one({"id": sid}, {"_id":0})
    if not s: raise HTTPException(404, "Student not found")
    return s

@router.get("/students/{sid}/eligible-receipt-types")
async def eligible_receipt_types(sid: str, user = Depends(get_current_user)):
    """Which receipt type(s) this student's actual class/medium/stream/bus
    status qualify for, using the SAME eligible_receipt_codes_for_class rule
    the backend enforces at receipt-creation time (core.py) — so the frontend
    never has its own, potentially-drifting copy of this logic.
    Returns: { primary, eligible: [{id,code,name}], bus_eligible, notes }
    `primary` is the single most-specific match (never EMJC) for auto-select;
    `eligible` lists every receipt_type this student could legitimately use,
    including EMJC (as a non-primary option) and BUS only if bus_required."""
    student = await db.students.find_one({"id": sid}, {"_id": 0})
    if not student:
        raise HTTPException(404, "Student not found")
    class_doc = await db.classes.find_one({"id": student.get("class_id")}, {"_id": 0}) if student.get("class_id") else None
    dept = await db.departments.find_one({"id": student.get("department_id")}, {"_id": 0}) if student.get("department_id") else None
    dept_code = dept.get("code") if dept else None

    specific_codes, notes = eligible_receipt_codes_for_class(
        class_doc.get("name") if class_doc else None,
        student.get("medium"),
        student.get("stream"),
    )
    all_types = await db.receipt_types.find({"enabled": True}, {"_id": 0}).to_list(50)

    eligible = []
    matched_types = []
    for t in all_types:
        code = t.get("code")
        allowed = t.get("applicable_dept_codes") or []
        if dept_code and allowed and dept_code not in allowed:
            continue  # wrong department entirely
        if code == "BUS":
            if student.get("bus_required"):
                eligible.append({"id": t["id"], "code": code, "name": t["name"]})
            continue
        if code == "DV":
            continue  # not a student academic receipt
        if code == "EMJC":
            eligible.append({"id": t["id"], "code": code, "name": t["name"]})  # always available, never primary
            continue
        if code in specific_codes:
            eligible.append({"id": t["id"], "code": code, "name": t["name"]})
            matched_types.append(t)

    # When more than one specific code matches (only possible for Class 11/12,
    # where both JC and JCACS can be eligible for a canonical stream), JC is
    # always the auto-selected primary — JCACS stays listed in `eligible` for
    # manual selection but is never auto-picked. For every other class this is
    # a no-op since there's only ever one matched type.
    primary = None
    if matched_types:
        jc_match = next((t for t in matched_types if t.get("code") == "JC"), None)
        primary = (jc_match or matched_types[0])["id"]

    return {
        "primary": primary,
        "eligible": eligible,
        "bus_eligible": bool(student.get("bus_required")),
        "class_name": class_doc.get("name") if class_doc else None,
        "medium": student.get("medium"),
        "stream": student.get("stream"),
        "notes": notes,
    }

@router.get("/students/{sid}/ledger")
async def student_ledger(sid: str, user = Depends(get_current_user)):
    s = await db.students.find_one({"id": sid}, {"_id":0})
    if not s: raise HTTPException(404, "Not found")
    # Denormalised purely for display (Student Profile academic-identity
    # header) - class_id/department_id stay the source of truth everywhere
    # else; medium/stream already live directly on the student document for
    # Junior College, class_name/department_name do not and are looked up
    # once here rather than pushed onto every frontend consumer.
    if s.get("class_id"):
        cls = await db.classes.find_one({"id": s["class_id"]}, {"_id": 0})
        if cls:
            s["class_name"] = cls.get("name")
            if not s.get("medium"): s["medium"] = cls.get("medium")
            if not s.get("stream"): s["stream"] = cls.get("stream")
    if s.get("department_id"):
        dept = await db.departments.find_one({"id": s["department_id"]}, {"_id": 0, "name": 1})
        if dept: s["department_name"] = dept.get("name")
    receipts = await db.receipts.find({"student_id": sid, "status":{"$ne":"cancelled"}}, {"_id":0}).sort("created_at", -1).to_list(500)
    adjustments = await db.adjustments.find({"student_id": sid, "status":"approved"}, {"_id":0}).to_list(200)
    fs = None
    if s.get("fee_structure_id"):
        fs = await db.fee_structures.find_one({"id": s["fee_structure_id"]}, {"_id":0})
    if s.get("bus_stop_no"):
        stop = await db.bus_stops.find_one({"stop_no": s["bus_stop_no"]}, {"_id":0})
        if stop:
            s["bus_stop_monthly_fee"] = stop.get("monthly_fee")
            s["bus_main_area"] = stop.get("main_area")
    bus_assignments = await db.bus_assignments.find({"student_id": sid}, {"_id":0}).sort("effective_from", -1).to_list(200)
    bus_charges = await db.bus_charges.find({"student_id": sid}, {"_id":0}).sort("month", -1).to_list(500)
    bus_outstanding = sum(max(0, c.get("amount", 0) - c.get("amount_paid", 0)) for c in bus_charges if c.get("status") != "paid")
    ob_doc = await db.student_opening_balances.find_one(
        {"student_id": sid, "academic_year": s.get("academic_year") or "2026-27"}, {"_id": 0}
    )
    opening_balance = ob_doc.get("amount", 0) if ob_doc else 0

    # Opening-paid — amount this student had already paid BEFORE FeeHub went
    # live (verified from the school's real prior records/Excel, imported into
    # the fee_details collection with a per-row audit trail — never a fake
    # receipt, per the explicit rule against synthetic transactions). fee_details
    # is otherwise deliberately decoupled from the live receipts ledger, but a
    # genuine pre-go-live payment must still count as "paid" on the Fee Update
    # screen, or every migrated student wrongly shows ₹0 paid despite real
    # money already collected. Added once, on top of live receipts — never
    # replacing them — so this is additive carry-forward, not double counting.
    fee_detail_doc = await db.fee_details.find_one(
        {"student_id": sid, "academic_year": s.get("academic_year") or "2026-27"}, {"_id": 0}
    )
    opening_paid = float(fee_detail_doc.get("total_paid") or 0) if fee_detail_doc else 0

    # Per-item breakdown (Option A) — merges this student's fee_overrides on top
    # of the shared fee_structure.items[], with paid/outstanding/status always
    # DERIVED live from real receipt lines (never a stored status field). This
    # is now the single authoritative source for "what does this student still
    # owe, per fee head/installment" — the frontend no longer computes this
    # itself. Existing students with no override see exactly the shared items,
    # unchanged from before this field was added.
    overrides = await db.student_fee_overrides.find(
        {"student_id": sid, "academic_year": s.get("academic_year") or "2026-27"}, {"_id": 0}
    ).to_list(20)
    fee_items = compute_fee_items((fs.get("items") if fs else None) or [], overrides, receipts)

    # Carry the opening-paid figure INTO the per-head breakdown too (not just
    # the top-level total) — filled head-by-head, oldest/first head in the fee
    # structure first, until exhausted. Without this, the header "Paid" card
    # would be correct but every individual head below it would still show
    # ₹0 paid / full outstanding, so New Receipt would prompt the cashier to
    # re-collect money the family already paid before FeeHub existed. Any
    # leftover that doesn't fit inside the student's current fee items (e.g.
    # the fee structure changed since the historical payment) is still added
    # into total_paid below, so the headline figure never silently loses it.
    fee_items, opening_paid_unabsorbed = apply_opening_paid(fee_items, opening_paid)

    # total_paid = SCHOOL fee paid only, matched per fee-head line exactly like
    # compute_fee_items() does — NOT a raw sum of every receipt's total. Bus Fee
    # and Bus receipts are tracked entirely separately (bus_charges/bus_outstanding
    # above); a bus payment must never silently reduce the school total below,
    # or "Outstanding" would drop by the bus amount even though no school fee
    # was paid. (Previously this summed every non-refund/non-debit_voucher
    # receipt's total regardless of type, which double-counted bus/misc
    # payments against the school fee balance — fixed here.) Opening-paid
    # (pre-FeeHub, imported from fee_details, see above) is already folded
    # into fee_items' paid; opening_paid_unabsorbed covers only the remainder
    # that didn't fit any current head.
    total_paid = sum(it["paid"] for it in fee_items) + opening_paid_unabsorbed
    total_refunded = sum(r.get("total", 0) for r in receipts if r.get("receipt_type") == "refund")
    total_adjusted = sum(a.get("amount", 0) for a in adjustments)
    school_payable = (fs.get("total") if fs else 0) - total_paid - total_adjusted + total_refunded

    return {
        "student": s, "fee_structure": fs, "receipts": receipts, "adjustments": adjustments,
        "bus_assignments": bus_assignments, "bus_charges": bus_charges, "bus_outstanding": bus_outstanding,
        "opening_paid": opening_paid, "fee_details_on_record": fee_detail_doc,
        "total_paid": total_paid, "total_refunded": total_refunded,
        "total_adjusted": total_adjusted, "school_outstanding": max(0, school_payable),
        "opening_balance": opening_balance,
        "outstanding": max(0, school_payable) + bus_outstanding + opening_balance,
        "fee_items": fee_items, "fee_overrides": overrides,
    }

@router.post("/students")
async def create_student(body: StudentIn, user = Depends(require_roles("administrator","manager","accountant","cashier"))):
    existing = await db.students.find_one({"admission_no": body.admission_no})
    if existing:
        raise HTTPException(400, "Admission number already exists")
    sid = gen_id()
    doc = {"id": sid, **body.model_dump(), "status":"active", "created_at": now_iso()}
    await db.students.insert_one(doc)
    await audit(user, "create", "student", sid, {"admission_no": body.admission_no})
    return {k:v for k,v in doc.items() if k != "_id"}

@router.post("/students/bulk-import")
async def bulk_import_students(body: Dict[str, Any], user = Depends(require_roles("administrator","manager","accountant"))):
    """Import students with MANDATORY Medium column and automatic fee-structure assignment.

    Required columns: admission_no, name, medium, class_name.
    JC rows must also include `stream`.
    Extra: father_name, mother_name, guardian_mobile, section, roll_no, academic_year, address.
    Rejects rows where Medium/class combination has no matching fee structure."""
    from core import canonical_medium, canonical_stream, normalize_class_name, resolve_fee_structure
    rows = body.get("rows", [])
    if not isinstance(rows, list) or not rows:
        raise HTTPException(400, "rows must be a non-empty array")
    batch_id = body.get("batch_id") or gen_id()
    depts_by_code = {d["code"]: d for d in await db.departments.find({}, {"_id":0}).to_list(100)}
    all_classes = await db.classes.find({}, {"_id":0}).to_list(1000)
    created, skipped, errors = 0, 0, []
    for idx, r in enumerate(rows):
        try:
            adm = str(r.get("admission_no","")).strip()
            name = str(r.get("name","")).strip()
            raw_medium = str(r.get("medium","")).strip()
            raw_class = str(r.get("class_name","")).strip()
            raw_stream = str(r.get("stream","")).strip()

            if not adm or not name:
                errors.append({"row": idx+1, "error": "admission_no and name are required", "data": r}); continue
            if not raw_medium:
                errors.append({"row": idx+1, "error": "Medium is required (English Medium / Semi Medium (Marathi) / Junior College)", "data": r}); continue
            medium = canonical_medium(raw_medium)
            if not medium:
                errors.append({"row": idx+1, "error": f"invalid Medium '{raw_medium}' — use English Medium, Semi Medium (Marathi), or Junior College", "data": r}); continue
            if not raw_class:
                errors.append({"row": idx+1, "error": "class_name is required", "data": r}); continue
            class_name = normalize_class_name(raw_class)
            stream = None
            if medium == "Junior College":
                if not raw_stream:
                    errors.append({"row": idx+1, "error": "Junior College students must include a Stream (Arts/Commerce/Science/Bi-Focal)", "data": r}); continue
                stream = canonical_stream(raw_stream)
                if not stream:
                    errors.append({"row": idx+1, "error": f"unknown stream '{raw_stream}' — allowed: Arts, Commerce, Science, Bi-Focal", "data": r}); continue
            elif raw_stream:
                # Non-JC row must NOT carry a stream — protects against Class 5 English being mis-tagged
                errors.append({"row": idx+1, "error": f"Stream '{raw_stream}' only allowed for Junior College rows", "data": r}); continue
            # Medium/class alignment sanity check
            if medium == "Junior College" and class_name not in ("Class 11", "Class 12"):
                errors.append({"row": idx+1, "error": f"Junior College only supports Class 11 / Class 12 — got '{class_name}'", "data": r}); continue
            if medium != "Junior College" and class_name in ("Class 11", "Class 12"):
                errors.append({"row": idx+1, "error": f"Class 11/12 must use Junior College medium — got '{medium}'", "data": r}); continue
            if await db.students.find_one({"admission_no": adm}):
                skipped += 1; continue

            # Pick department to match the seed_2026 logic - Secondary (Class 9/10) covers
            # BOTH mediums, same as catalog.py's pick_dept. A Marathi Semi 9th/10th student
            # must land in Secondary, not Marathi Primary.
            if medium == "Junior College":
                dept = depts_by_code.get("JC")
            elif class_name in ("Class 9", "Class 10"):
                dept = depts_by_code.get("SEC")
            elif medium == "English Medium":
                dept = depts_by_code.get("EP")
            else:
                dept = depts_by_code.get("MP")
            if not dept:
                errors.append({"row": idx+1, "error": f"department not configured for medium '{medium}'", "data": r}); continue

            # Find or auto-create the class row for this dept/medium/name
            class_query: Dict[str, Any] = {"department_id": dept["id"], "name": class_name, "medium": medium}
            if stream: class_query["stream"] = stream
            cls = next((c for c in all_classes
                        if c["department_id"] == dept["id"]
                        and c.get("name","").lower() == class_name.lower()
                        and (c.get("medium") or medium) == medium
                        and (stream is None or c.get("stream") == stream)), None)
            if not cls:
                cls = {"id": gen_id(), **class_query, "created_at": now_iso()}
                await db.classes.insert_one(cls); all_classes.append(cls)

            first_year_in_college = str(r.get("first_year_in_college","")).strip().lower() in ("y","yes","true","1","new")
            # Optional bus stop assignment (denormalised for receipts)
            bus_stop_no = None
            bus_stop_name = None
            bus_stop_raw = str(r.get("bus_stop_no","")).strip()
            if bus_stop_raw and bus_stop_raw.lower() not in ("nan","none","0"):
                try:
                    bus_stop_no = int(float(bus_stop_raw))
                except Exception:
                    errors.append({"row": idx+1, "error": f"bus_stop_no '{bus_stop_raw}' is not a number", "data": r}); continue
                stop = await db.bus_stops.find_one({"stop_no": bus_stop_no}, {"_id": 0})
                if not stop:
                    errors.append({"row": idx+1, "error": f"bus_stop_no {bus_stop_no} not found in master list — seed the 2026-27 bus stops first", "data": r}); continue
                bus_stop_name = stop.get("stop_name")

            fs = await resolve_fee_structure(medium, class_name, stream,
                                              first_year_in_college=first_year_in_college)
            if not fs:
                errors.append({"row": idx+1, "error": f"No approved fee structure for {medium} · {class_name}" + (f" · {stream}" if stream else "") + ". Ask admin to seed the 2026-27 structure first.", "data": r}); continue

            await db.students.insert_one({
                "id": gen_id(),
                "admission_no": adm, "name": name,
                "father_name": r.get("father_name"), "mother_name": r.get("mother_name"),
                "guardian_name": r.get("guardian_name") or r.get("father_name") or r.get("mother_name"),
                "guardian_mobile": str(r.get("guardian_mobile","") or r.get("mobile_number","") or "").strip(),
                "section": (r.get("section") or "").strip() or None,
                "roll_no":  (r.get("roll_no")  or r.get("roll_number") or "").strip() or None,
                "medium": medium, "stream": stream,
                "first_year_in_college": first_year_in_college,
                "bus_stop_no": bus_stop_no, "bus_stop_name": bus_stop_name,
                "department_id": dept["id"], "class_id": cls["id"],
                "fee_structure_id": fs["id"],
                "academic_year": (r.get("academic_year") or "2026-27").strip(),
                "address": r.get("address"), "status":"active", "created_at": now_iso(),
                "imported_at": now_iso(), "imported_by": user["id"], "import_batch_id": batch_id,
            })
            created += 1
        except Exception as e:
            errors.append({"row": idx+1, "error": str(e), "data": r})
    await db.import_batches.insert_one({
        "id": batch_id, "type": "students", "created": created, "skipped": skipped,
        "errors_count": len(errors), "total": len(rows),
        "user_id": user["id"], "user_name": user["name"], "created_at": now_iso(),
    })
    await audit(user, "bulk_import", "student", batch_id, {"created": created, "skipped": skipped, "errors": len(errors)})
    return {"created": created, "skipped": skipped, "errors": errors, "total": len(rows), "batch_id": batch_id}

@router.post("/students/bulk-delete")
async def bulk_delete_students(body: Dict[str, Any], user = Depends(require_roles("administrator","manager"))):
    batch_id = body.get("batch_id")
    ids = body.get("student_ids") or []
    q: Dict[str, Any] = {}
    if batch_id:
        q["import_batch_id"] = batch_id
    elif ids:
        q["id"] = {"$in": ids}
    else:
        raise HTTPException(400, "Provide batch_id or student_ids")
    stus = await db.students.find(q, {"_id": 0}).to_list(5000)
    if not stus:
        return {"deleted": 0, "protected_with_receipts": 0}
    with_receipts = await db.receipts.distinct("student_id", {"student_id": {"$in": [s["id"] for s in stus]}})
    protected_ids = set(with_receipts)
    deletable = [s["id"] for s in stus if s["id"] not in protected_ids]
    if deletable:
        await db.students.delete_many({"id": {"$in": deletable}})
    if batch_id:
        await db.import_batches.update_one({"id": batch_id}, {"$set": {"undone_at": now_iso(), "undone_by": user["name"], "undone_deleted": len(deletable), "undone_protected": len(protected_ids)}})
    await audit(user, "bulk_delete", "student", batch_id or "", {"deleted": len(deletable), "protected": len(protected_ids)})
    return {"deleted": len(deletable), "protected_with_receipts": len(protected_ids), "batch_id": batch_id}

@router.get("/students/{sid}/siblings")
async def student_siblings(sid: str, user = Depends(get_current_user)):
    s = await db.students.find_one({"id": sid}, {"_id":0})
    if not s: raise HTTPException(404, "Not found")
    gm = (s.get("guardian_mobile") or "").strip()
    if not gm: return {"siblings": []}
    others = await db.students.find(
        {"guardian_mobile": gm, "status": "active", "id": {"$ne": sid}},
        {"_id":0}
    ).to_list(20)
    return {"siblings": others}

@router.post("/students/bulk-reassign")
async def bulk_reassign_students(body: Dict[str, Any], user = Depends(require_roles("administrator","manager","accountant"))):
    ids = body.get("student_ids", [])
    to_class_id = body.get("to_class_id")
    if not ids or not to_class_id:
        raise HTTPException(400, "student_ids and to_class_id required")
    to_cls = await db.classes.find_one({"id": to_class_id})
    if not to_cls: raise HTTPException(400, "Target class not found")
    upd: Dict[str, Any] = {"class_id": to_class_id, "department_id": to_cls["department_id"]}
    if body.get("to_fee_structure_id"): upd["fee_structure_id"] = body["to_fee_structure_id"]
    result = await db.students.update_many({"id": {"$in": ids}}, {"$set": upd, "$push": {"reassign_history": {"to_class_id": to_class_id, "at": now_iso(), "by": user["name"]}}})
    await audit(user, "bulk_reassign", "student", "", {"count": result.modified_count, "to_class_id": to_class_id})
    return {"reassigned": result.modified_count}

@router.patch("/students/{sid}")
async def update_student(sid: str, body: Dict[str,Any], user = Depends(require_roles("administrator","manager","accountant"))):
    upd = {k:v for k,v in body.items() if k in ("name","class_id","section","guardian_name","guardian_mobile","address","fee_structure_id","bus_route","status")}
    await db.students.update_one({"id": sid}, {"$set": upd})
    await audit(user, "update", "student", sid, upd)
    return {"ok": True}

# ---------- Bus assignment (with history — see PART 31/32) ----------
# The exact fee is captured on the assignment record at the moment it's
# created, from the bus_stops master list for the given academic_year.
# Later fare changes on the master stop (bus-stops PATCH / bulk-update)
# never rewrite this captured value, so a student's historical bus charge
# stays exactly what it was when they were actually billed.
@router.post("/students/{sid}/bus-assignment")
async def assign_bus_stop(sid: str, body: Dict[str, Any], user = Depends(require_roles("administrator","manager","accountant"))):
    student = await db.students.find_one({"id": sid})
    if not student: raise HTTPException(404, "Student not found")
    try:
        stop_no = int(body.get("stop_no"))
    except Exception:
        raise HTTPException(400, "stop_no is required")
    academic_year = body.get("academic_year") or "2026-27"
    stop = await db.bus_stops.find_one({"stop_no": stop_no, "academic_year": academic_year}, {"_id": 0})
    if not stop:
        stop = await db.bus_stops.find_one({"stop_no": stop_no}, {"_id": 0})
    if not stop:
        raise HTTPException(404, f"Bus stop #{stop_no} not found")
    effective_from = body.get("effective_from") or now_iso()[:10]
    # Optional per-student fee override (Phase 6/7 "individual update") - when
    # given, THIS student's new assignment captures this explicit fee instead
    # of the stop's standard rate (e.g. a sibling discount). Omit it (every
    # existing caller does) and behavior is byte-for-byte unchanged - the
    # stop's own current monthly_fee is captured, exactly as before.
    fee_override = body.get("monthly_fee_override")
    reason = body.get("reason")
    effective_fee = float(fee_override) if fee_override is not None else stop.get("monthly_fee", 0)

    # Close any currently-active assignment (history is preserved, never overwritten).
    prior = await db.bus_assignments.find_one({"student_id": sid, "status": "active"}, {"_id": 0})
    await db.bus_assignments.update_many(
        {"student_id": sid, "status": "active"},
        {"$set": {"status": "inactive", "effective_to": effective_from}},
    )
    doc = {
        "id": gen_id(), "student_id": sid,
        "stop_no": stop["stop_no"], "main_area": stop.get("main_area", ""), "stop_name": stop.get("stop_name", ""),
        "monthly_fee": effective_fee, "academic_year": stop.get("academic_year", academic_year),
        "effective_from": effective_from, "effective_to": None, "status": "active",
        "created_at": now_iso(), "created_by": user["name"],
    }
    await db.bus_assignments.insert_one(doc)
    if fee_override is not None:
        await audit(user, "individual_fee_update", "bus_assignment", doc["id"], {
            "student_id": sid, "student_name": student.get("name"), "admission_no": student.get("admission_no"),
            "main_area": stop.get("main_area"), "stop_name": stop.get("stop_name"), "stop_no": stop["stop_no"],
            "old_monthly_fee": prior.get("monthly_fee") if prior else stop.get("monthly_fee", 0),
            "new_monthly_fee": effective_fee, "updated_by": user["name"], "updated_by_id": user["id"],
            "at": now_iso(), "reason": reason or "",
        })
    # Denormalized snapshot on the student doc for fast reads (list views, receipts, reports).
    await db.students.update_one({"id": sid}, {"$set": {
        "bus_required": True, "bus_stop_no": stop["stop_no"], "bus_stop_name": stop.get("stop_name", ""),
        "bus_main_area": stop.get("main_area", ""),
    }})
    await audit(user, "assign_bus_stop", "student", sid, {"stop_no": stop_no, "main_area": stop.get("main_area"), "stop_name": stop.get("stop_name")})
    return {k: v for k, v in doc.items() if k != "_id"}

@router.post("/students/{sid}/bus-assignment/remove")
async def remove_bus_assignment(sid: str, user = Depends(require_roles("administrator","manager","accountant"))):
    student = await db.students.find_one({"id": sid})
    if not student: raise HTTPException(404, "Student not found")
    today = now_iso()[:10]
    await db.bus_assignments.update_many(
        {"student_id": sid, "status": "active"},
        {"$set": {"status": "inactive", "effective_to": today}},
    )
    await db.students.update_one({"id": sid}, {"$set": {
        "bus_required": False, "bus_stop_no": None, "bus_stop_name": None, "bus_main_area": None,
    }})
    await audit(user, "remove_bus_assignment", "student", sid, {})
    return {"ok": True}

@router.get("/students/{sid}/bus-assignments")
async def list_bus_assignments(sid: str, user = Depends(get_current_user)):
    return await db.bus_assignments.find({"student_id": sid}, {"_id": 0}).sort("effective_from", -1).to_list(200)

# ---------- Bus Assignment data-entry XLSX/CSV (real dropdowns from the live Bus Stop Master) ----------

def _sanitize_named_range(raw: str, prefix: str = "MS") -> str:
    import re
    s = re.sub(r"[^A-Za-z0-9]+", "_", raw).strip("_")
    if not s or s[0].isdigit():
        s = "S_" + s
    return f"{prefix}_{s}"[:250]

@router.get("/bus-assignment-template.xlsx")
async def bus_assignment_template_xlsx(academic_year: str = "2026-27", only_unassigned: bool = False,
                                        user = Depends(require_roles("administrator","manager","accountant"))):
    """Real cascading Main Stop -> Sub Stop dropdowns built from the CURRENT Bus Stop Master
    at request time (never a hardcoded or stale list) - the Excel technique is a per-row
    INDIRECT(VLOOKUP(...)) so the Sub Stop dropdown always matches whatever Main Stop that row
    actually has selected, with no risk of the sanitized-name lookup drifting out of sync."""
    import io
    from collections import defaultdict
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    classes = {c["id"]: c for c in await db.classes.find({}, {"_id": 0}).to_list(2000)}
    stops = await db.bus_stops.find({"active": {"$ne": False}}, {"_id": 0}).sort([("main_area", 1), ("stop_no", 1)]).to_list(500)
    by_area = defaultdict(list)
    for st in stops:
        by_area[st["main_area"]].append(st)
    main_stops = sorted(by_area.keys())
    name_map: Dict[str, str] = {}
    used = set()
    for area in main_stops:
        nm = _sanitize_named_range(area); base = nm; n = 2
        while nm in used:
            nm = f"{base}_{n}"; n += 1
        used.add(nm); name_map[area] = nm

    sq: Dict[str, Any] = {"status": "active"}
    if only_unassigned:
        sq["bus_required"] = {"$ne": True}
    students = await db.students.find(sq, {"_id": 0}).sort("name", 1).to_list(5000)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bus Assignment"
    ws.append(["Balaji Convent & Junior College - Bus Student Assignment"])
    ws.append([f"Academic Year: {academic_year}", "", "", "", "", "", "", f"{len(stops)} bus stops across {len(main_stops)} Main Stops"])
    ws.append([])
    HEADERS = ["Admission No.", "Student Name", "Class", "Medium", "Bus Required (Yes/No)", "Main Stop", "Sub Stop", "Monthly Bus Fee (auto)"]
    header_row = 4
    ws.append(HEADERS)
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1E40AF")
        cell.alignment = Alignment(horizontal="center")

    LAST_DATA_ROW = header_row + max(len(students), 50)
    for i, s in enumerate(students):
        row = header_row + 1 + i
        cls = classes.get(s.get("class_id"), {})
        ws.cell(row=row, column=1, value=s["admission_no"])
        ws.cell(row=row, column=2, value=s["name"])
        ws.cell(row=row, column=3, value=cls.get("name", ""))
        ws.cell(row=row, column=4, value=s.get("medium", ""))
        ws.cell(row=row, column=5, value="Yes" if s.get("bus_required") else "No")
        ws.cell(row=row, column=6, value=s.get("bus_main_area") or "")
        ws.cell(row=row, column=7, value=s.get("bus_stop_name") or "")
    for row in range(header_row + 1, LAST_DATA_ROW + 1):
        ws.cell(row=row, column=8, value=f'=IFERROR(SUMIFS(BusStopMaster!$C:$C,BusStopMaster!$A:$A,F{row},BusStopMaster!$B:$B,G{row}),"")')

    for idx, w in enumerate([14, 24, 10, 22, 20, 26, 34, 20], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = f"A{header_row+1}"

    ms = wb.create_sheet("BusStopMaster")
    ms.append(["Main Stop", "Sub Stop", "Monthly Fee", "Stop No.", "Active"])
    for cell in ms[1]: cell.font = Font(bold=True)
    row_i = 2
    area_row_ranges = {}
    for area in main_stops:
        start_row = row_i
        for st in sorted(by_area[area], key=lambda x: x["stop_no"]):
            ms.cell(row=row_i, column=1, value=area)
            ms.cell(row=row_i, column=2, value=st["stop_name"])
            ms.cell(row=row_i, column=3, value=st["monthly_fee"])
            ms.cell(row=row_i, column=4, value=st["stop_no"])
            ms.cell(row=row_i, column=5, value="Yes")
            row_i += 1
        area_row_ranges[area] = (start_row, row_i - 1)
    for col, w in zip("ABCDE", [26, 34, 12, 10, 8]):
        ms.column_dimensions[col].width = w

    ls = wb.create_sheet("Lists")
    ls.append(["MainStops", "SanitizedKey"])
    for i, area in enumerate(main_stops, start=2):
        ls.cell(row=i, column=1, value=area)
        ls.cell(row=i, column=2, value=name_map[area])
    ls.column_dimensions["A"].width = 26
    ls.column_dimensions["B"].width = 30

    wb.defined_names["MainStopList"] = DefinedName("MainStopList", attr_text=f"Lists!$A$2:$A${1+len(main_stops)}")
    for area in main_stops:
        r1, r2 = area_row_ranges[area]
        wb.defined_names[name_map[area]] = DefinedName(name_map[area], attr_text=f"BusStopMaster!$B${r1}:$B${r2}")

    dv_main = DataValidation(type="list", formula1="=MainStopList", allow_blank=True, showErrorMessage=True,
                              errorTitle="Invalid Main Stop", error="Pick a Main Stop from the FeeHub Bus Stop Master list.")
    ws.add_data_validation(dv_main)
    dv_main.add(f"F{header_row+1}:F{LAST_DATA_ROW}")
    for r in range(header_row + 1, LAST_DATA_ROW + 1):
        dv_row = DataValidation(
            type="list", formula1=f'=INDIRECT(VLOOKUP(F{r},Lists!$A$2:$B${1+len(main_stops)},2,FALSE))',
            allow_blank=True, showErrorMessage=True,
            errorTitle="Invalid Sub Stop", error="Pick a Sub Stop that belongs to the selected Main Stop.",
        )
        ws.add_data_validation(dv_row)
        dv_row.add(f"G{r}")

    buf = io.BytesIO()
    wb.save(buf)
    await audit(user, "download_template", "bus_assignment", "", {"academic_year": academic_year, "student_count": len(students)})
    return Response(content=buf.getvalue(),
                     media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     headers={"Content-Disposition": f'attachment; filename="Bus_Assignment_Template_{academic_year}.xlsx"'})

@router.get("/students/export.csv")
async def export_students_csv(
    department_id: Optional[str] = None,
    academic_year: Optional[str] = None,
    user = Depends(require_roles("administrator","manager","accountant")),
):
    """Student export for receipt-eligibility purposes (Task 7/8). Column
    order and the Receipt Type column are built using the SAME
    eligible_receipt_codes_for_class rule the backend enforces at receipt
    creation (core.py) and the /students/{id}/eligible-receipt-types endpoint
    use — so this file can never state a different rule than the app itself
    enforces. Read-only: does not write anything. All values are pulled
    directly from existing student/class/department records — nothing here
    is invented or guessed."""
    import io, csv as _csv
    classes = {c["id"]: c for c in await db.classes.find({}, {"_id": 0}).to_list(2000)}
    departments = {d["id"]: d for d in await db.departments.find({}, {"_id": 0}).to_list(50)}
    receipt_types = await db.receipt_types.find({"enabled": True}, {"_id": 0}).to_list(50)

    sq: Dict[str, Any] = {"status": "active"}
    if department_id: sq["department_id"] = department_id
    if academic_year: sq["academic_year"] = academic_year
    students = await db.students.find(sq, {"_id": 0}).sort("name", 1).to_list(10000)

    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow([
        "Admission No.", "Student Name", "Class", "Division", "Medium", "Stream",
        "Academic Year", "Department", "Bus Facility Status", "Bus Area/Village",
        "Bus Sub Stop", "Father Name", "Mother Name", "Contact No.",
        "Eligible Receipt Type(s)", "Primary Receipt Type", "Eligibility Note",
    ])
    for s in students:
        cls = classes.get(s.get("class_id"), {})
        dept = departments.get(s.get("department_id"), {})
        specific_codes, notes = eligible_receipt_codes_for_class(cls.get("name"), s.get("medium"), s.get("stream"))
        dept_code = dept.get("code")
        codes_for_row = []
        for t in receipt_types:
            code = t.get("code")
            allowed = t.get("applicable_dept_codes") or []
            if dept_code and allowed and dept_code not in allowed:
                continue
            if code == "BUS":
                if s.get("bus_required"): codes_for_row.append(code)
            elif code == "DV":
                continue
            elif code == "EMJC":
                codes_for_row.append(code)
            elif code in specific_codes:
                codes_for_row.append(code)
        primary = next((c for c in codes_for_row if c != "EMJC" and c != "BUS"), (codes_for_row[0] if codes_for_row else ""))
        w.writerow([
            s.get("admission_no",""), s.get("name",""), cls.get("name",""), s.get("section",""),
            s.get("medium",""), s.get("stream",""), s.get("academic_year",""), dept.get("name",""),
            "Yes" if s.get("bus_required") else "No", s.get("bus_main_area") or "", s.get("bus_stop_name") or "",
            s.get("father_name",""), s.get("mother_name",""), s.get("guardian_mobile",""),
            "; ".join(codes_for_row), primary, "; ".join(notes),
        ])
    return Response(content=buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                     headers={"Content-Disposition": 'attachment; filename="FeeHub_Students_Export.csv"'})

@router.get("/bus-assignment-template.csv")
async def bus_assignment_template_csv(only_unassigned: bool = False,
                                       user = Depends(require_roles("administrator","manager","accountant"))):
    import io, csv as _csv
    classes = {c["id"]: c for c in await db.classes.find({}, {"_id": 0}).to_list(2000)}
    sq: Dict[str, Any] = {"status": "active"}
    if only_unassigned:
        sq["bus_required"] = {"$ne": True}
    students = await db.students.find(sq, {"_id": 0}).sort("name", 1).to_list(5000)
    buf = io.StringIO()
    w = _csv.writer(buf)
    w.writerow(["Admission No.", "Student Name", "Class", "Medium", "Bus Required", "Main Stop", "Sub Stop", "Monthly Bus Fee"])
    for s in students:
        cls = classes.get(s.get("class_id"), {})
        w.writerow([s["admission_no"], s["name"], cls.get("name",""), s.get("medium",""),
                    "Yes" if s.get("bus_required") else "No", s.get("bus_main_area") or "", s.get("bus_stop_name") or "", ""])
    return Response(content=buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                     headers={"Content-Disposition": 'attachment; filename="Bus_Assignment_Template.csv"'})

@router.post("/bus-assignment/bulk-import")
async def bulk_import_bus_assignments(body: Dict[str, Any],
                                       user = Depends(require_roles("administrator","manager","accountant"))):
    """Import bus assignments from an uploaded CSV/XLSX (already parsed to rows client-side).
    Matches students by admission_no ONLY. Every Main Stop + Sub Stop pair is validated against
    the CURRENT Bus Stop Master - an unrecognised stop is rejected, never silently created from
    a typo. Pass `preview: true` to validate/classify without writing anything."""
    rows: List[Dict[str, Any]] = body.get("rows", [])
    if not isinstance(rows, list) or not rows:
        raise HTTPException(400, "rows must be a non-empty array")
    preview = bool(body.get("preview"))
    academic_year = body.get("academic_year") or "2026-27"

    stops = await db.bus_stops.find({"active": {"$ne": False}}, {"_id": 0}).to_list(500)
    stop_by_pair = {(s["main_area"].strip().lower(), s["stop_name"].strip().lower()): s for s in stops}

    valid, invalid = [], []
    seen_adm = {}
    for idx, r in enumerate(rows):
        row_no = idx + 1
        adm = str(r.get("admission_no", "")).strip()
        bus_required_raw = str(r.get("bus_required", "")).strip().lower()
        main_stop = str(r.get("main_stop", "")).strip()
        sub_stop = str(r.get("sub_stop", "")).strip()
        if not adm:
            invalid.append({"row": row_no, "error": "admission_no is required", "data": r}); continue
        if adm in seen_adm:
            invalid.append({"row": row_no, "error": f"admission_no '{adm}' also appears at row {seen_adm[adm]} in this file", "data": r}); continue
        seen_adm[adm] = row_no
        student = await db.students.find_one({"admission_no": adm}, {"_id": 0})
        if not student:
            invalid.append({"row": row_no, "error": f"No student found with admission_no '{adm}'", "data": r}); continue
        bus_required = bus_required_raw in ("yes", "y", "true", "1")
        if not bus_required:
            valid.append({"row": row_no, "admission_no": adm, "student_id": student["id"], "action": "remove", "stop": None})
            continue
        if not main_stop or not sub_stop:
            invalid.append({"row": row_no, "error": f"Bus Required is Yes but Main Stop / Sub Stop is missing for '{adm}'", "data": r}); continue
        stop = stop_by_pair.get((main_stop.strip().lower(), sub_stop.strip().lower()))
        if not stop:
            invalid.append({"row": row_no, "error": f"'{main_stop}' / '{sub_stop}' is not a stop configured in the Bus Stop Master — not imported (no new stop was created)", "data": r}); continue
        valid.append({"row": row_no, "admission_no": adm, "student_id": student["id"], "action": "assign", "stop": stop})

    result = {
        "total_rows": len(rows), "valid_rows": len(valid), "invalid_rows": len(invalid),
        "rows_to_assign": sum(1 for v in valid if v["action"] == "assign"),
        "rows_to_remove": sum(1 for v in valid if v["action"] == "remove"),
        "valid": valid, "invalid": invalid,
    }
    if preview:
        return {**result, "committed": False}

    assigned = removed = 0
    for v in valid:
        sid = v["student_id"]
        today = now_iso()[:10]
        if v["action"] == "remove":
            await db.bus_assignments.update_many({"student_id": sid, "status": "active"}, {"$set": {"status": "inactive", "effective_to": today}})
            await db.students.update_one({"id": sid}, {"$set": {"bus_required": False, "bus_stop_no": None, "bus_stop_name": None, "bus_main_area": None}})
            removed += 1
        else:
            stop = v["stop"]
            await db.bus_assignments.update_many({"student_id": sid, "status": "active"}, {"$set": {"status": "inactive", "effective_to": today}})
            await db.bus_assignments.insert_one({
                "id": gen_id(), "student_id": sid, "stop_no": stop["stop_no"], "main_area": stop["main_area"],
                "stop_name": stop["stop_name"], "monthly_fee": stop.get("monthly_fee", 0),
                "academic_year": stop.get("academic_year", academic_year),
                "effective_from": today, "effective_to": None, "status": "active",
                "created_at": now_iso(), "created_by": user["name"],
            })
            await db.students.update_one({"id": sid}, {"$set": {
                "bus_required": True, "bus_stop_no": stop["stop_no"], "bus_stop_name": stop["stop_name"], "bus_main_area": stop["main_area"],
            }})
            assigned += 1
    await audit(user, "bulk_import", "bus_assignment", "", {"assigned": assigned, "removed": removed, "invalid": len(invalid)})
    return {"assigned": assigned, "removed": removed, "errors": invalid, "total": len(rows), "committed": True}

# ---------- Monthly Bus Fee generation ----------
# A charge is a DUE amount, never a receipt - it only affects "outstanding".
# Idempotency key: (student_id, academic_year, month, charge_type) is unique
# by construction (checked before every insert) so running generation twice,
# or running both the manual button and the automatic monthly job, can never
# double-charge - the second attempt always finds the existing charge and
# skips it. The charge amount is copied from the bus_assignment record that
# was active AT THE START of the target month (not a live lookup against the
# current bus_stops fee), so a later fare change never rewrites what a
# student was actually charged for a past month.
CHARGE_TYPE = "bus_monthly"
BUS_CHARGING_STATE_ID = "global"

async def _bus_charging_active() -> bool:
    """Global switch, separate from any individual student's bus assignment.
    Defaults to active (True) if never toggled - a fresh install charges
    normally without an admin having to explicitly turn it on first."""
    doc = await db.bus_charging_state.find_one({"id": BUS_CHARGING_STATE_ID})
    return doc.get("active", True) if doc else True

async def _find_assignment_for_month(sid: str, month_start: str) -> Optional[dict]:
    """The bus_assignments row effective for `month_start` ('YYYY-MM-01') -
    i.e. effective_from <= month_start < effective_to (or effective_to is
    still open). Assignment date strings are ISO 'YYYY-MM-DD', so plain
    string comparison is chronological."""
    return await db.bus_assignments.find_one({
        "student_id": sid,
        "effective_from": {"$lte": month_start},
        "$or": [{"effective_to": None}, {"effective_to": {"$gt": month_start}}],
    }, {"_id": 0}, sort=[("effective_from", -1)])

async def _bus_generation_run(academic_year: str, month: str, commit: bool, user: dict) -> Dict[str, Any]:
    if not month or len(month) != 7 or month[4] != "-":
        raise HTTPException(400, "month must be 'YYYY-MM', e.g. 2026-09")
    charging_active = await _bus_charging_active()
    if commit and not charging_active:
        # Preview is still allowed while stopped (read-only, useful to see what
        # WOULD be generated) - only the actual commit is blocked. The automatic
        # scheduler calls this with commit=True too, so this one check covers
        # both the manual button and the background job with zero duplication.
        return {
            "preview": False, "academic_year": academic_year, "month": month, "month_label": _month_label(month),
            "generated": 0, "already_existing": 0, "failed": 0, "total_amount": 0, "rows": [], "errors": [],
            "blocked": True, "blocked_reason": "Bus fee charging is currently STOPPED globally. Resume it from Admin -> Bus Fees before generating new charges.",
        }
    month_start = f"{month}-01"

    students = await db.students.find({"status": "active"}, {"_id": 0}).to_list(20000)
    to_generate, already, errors = [], [], []
    total_new = 0.0

    for s in students:
        assignment = await _find_assignment_for_month(s["id"], month_start)
        if not assignment:
            continue  # never had a bus stop, or stopped riding before this month
        existing = await db.bus_charges.find_one({
            "student_id": s["id"], "academic_year": academic_year, "month": month, "charge_type": CHARGE_TYPE,
        })
        if existing:
            already.append({"student_id": s["id"], "name": s["name"], "admission_no": s["admission_no"]})
            continue
        row = {
            "student_id": s["id"], "name": s["name"], "admission_no": s["admission_no"],
            "main_area": assignment["main_area"], "stop_name": assignment["stop_name"], "stop_no": assignment["stop_no"],
            "amount": assignment["monthly_fee"],
        }
        to_generate.append(row)
        total_new += float(assignment["monthly_fee"])

    if not commit:
        return {
            "preview": True, "academic_year": academic_year, "month": month, "month_label": _month_label(month),
            "active_bus_students": len(to_generate) + len(already),
            "already_generated": len(already), "to_generate": len(to_generate),
            "total_new_amount": total_new, "rows": to_generate, "errors": errors,
        }

    generated = []
    for row in to_generate:
        cid = gen_id()
        try:
            await db.bus_charges.insert_one({
                "id": cid, "student_id": row["student_id"],
                "main_area": row["main_area"], "stop_name": row["stop_name"], "stop_no": row["stop_no"],
                "amount": row["amount"], "amount_paid": 0, "status": "unpaid",
                "academic_year": academic_year, "month": month, "month_label": _month_label(month),
                "charge_type": CHARGE_TYPE,
                "generated_at": now_iso(), "generated_by": user["name"] if user else "system",
            })
            generated.append(row)
        except Exception as e:
            errors.append({"student_id": row["student_id"], "name": row["name"], "error": str(e)})

    if user:
        await audit(user, "generate_bus_charges", "bus_charge", "", {
            "academic_year": academic_year, "month": month,
            "generated": len(generated), "already_existing": len(already), "failed": len(errors),
        })
    return {
        "preview": False, "academic_year": academic_year, "month": month, "month_label": _month_label(month),
        "generated": len(generated), "already_existing": len(already), "failed": len(errors),
        "total_amount": sum(r["amount"] for r in generated), "rows": generated, "errors": errors,
    }

@router.post("/bus-charges/preview")
async def preview_bus_charges(body: Dict[str, Any], user = Depends(require_roles("administrator","manager"))):
    return await _bus_generation_run(body.get("academic_year") or "2026-27", body.get("month"), commit=False, user=user)

@router.post("/bus-charges/generate")
async def generate_bus_charges(body: Dict[str, Any], user = Depends(require_roles("administrator","manager"))):
    return await _bus_generation_run(body.get("academic_year") or "2026-27", body.get("month"), commit=True, user=user)

# ---------- Global "Stop/Resume Bus Fee Charging" (admin-only) ----------
# This is DELIBERATELY separate from an individual student's bus assignment
# (POST /students/{sid}/bus-assignment/remove). This one only gates whether
# NEW monthly charges get generated for ANYONE - it never touches existing
# charges, never touches any student's bus assignment, never touches the
# bus_stops master list, and never affects school fees.
@router.get("/bus-charges/status")
async def bus_charging_status(user = Depends(get_current_user)):
    doc = await db.bus_charging_state.find_one({"id": BUS_CHARGING_STATE_ID}, {"_id": 0})
    return doc or {"id": BUS_CHARGING_STATE_ID, "active": True, "changed_at": None, "changed_by": None, "reason": None}

@router.post("/bus-charges/stop")
async def stop_bus_charging(body: Dict[str, Any], user = Depends(require_roles("administrator"))):
    reason = (body.get("reason") or "").strip()
    if not reason:
        raise HTTPException(400, "A reason is required to stop bus fee charging.")
    prev = await _bus_charging_active()
    await db.bus_charging_state.update_one(
        {"id": BUS_CHARGING_STATE_ID},
        {"$set": {"active": False, "changed_at": now_iso(), "changed_by": user["name"], "reason": reason}},
        upsert=True,
    )
    await audit(user, "stop_bus_charging", "bus_charging_state", BUS_CHARGING_STATE_ID, {
        "previous_status": "ACTIVE" if prev else "STOPPED", "new_status": "STOPPED", "reason": reason,
    })
    return {"ok": True, "active": False}

@router.post("/bus-charges/resume")
async def resume_bus_charging(body: Dict[str, Any], user = Depends(require_roles("administrator"))):
    reason = (body.get("reason") or "").strip()
    prev = await _bus_charging_active()
    await db.bus_charging_state.update_one(
        {"id": BUS_CHARGING_STATE_ID},
        {"$set": {"active": True, "changed_at": now_iso(), "changed_by": user["name"], "reason": reason or None}},
        upsert=True,
    )
    await audit(user, "resume_bus_charging", "bus_charging_state", BUS_CHARGING_STATE_ID, {
        "previous_status": "ACTIVE" if prev else "STOPPED", "new_status": "ACTIVE", "reason": reason,
    })
    return {"ok": True, "active": True}

@router.get("/bus-charges/dashboard")
async def bus_charges_dashboard(user = Depends(get_current_user)):
    charging_active = await _bus_charging_active()
    active_bus_students = await db.students.count_documents({"status": "active", "bus_required": True})
    current_month = now_iso()[:7]
    prev_year, prev_mon = divmod(int(current_month[:4]) * 12 + int(current_month[5:7]) - 2, 12)
    previous_month = f"{prev_year}-{prev_mon + 1:02d}"

    current_month_charges = await db.bus_charges.find({"month": current_month}, {"_id": 0}).to_list(20000)
    current_month_total = sum(c["amount"] for c in current_month_charges)

    unpaid = await db.bus_charges.find({"status": {"$in": ["unpaid", "partial"]}}, {"_id": 0}).to_list(50000)
    total_bus_outstanding = sum(max(0, c["amount"] - c.get("amount_paid", 0)) for c in unpaid)
    previous_month_outstanding = sum(
        max(0, c["amount"] - c.get("amount_paid", 0)) for c in unpaid if c["month"] == previous_month
    )

    return {
        "charging_active": charging_active,
        "active_bus_students": active_bus_students,
        "current_month": current_month, "current_month_label": _month_label(current_month),
        "current_month_charges_count": len(current_month_charges), "current_month_charges_total": current_month_total,
        "previous_month": previous_month, "previous_month_label": _month_label(previous_month),
        "previous_month_outstanding": previous_month_outstanding,
        "total_bus_outstanding": total_bus_outstanding,
    }

@router.get("/bus-charges")
async def list_bus_charges(student_id: Optional[str] = None, month: Optional[str] = None, user = Depends(get_current_user)):
    q: Dict[str, Any] = {}
    if student_id: q["student_id"] = student_id
    if month: q["month"] = month
    return await db.bus_charges.find(q, {"_id": 0}).sort("month", -1).to_list(5000)

# ---------- Optional automatic monthly generation ----------
# Runs once a day and generates the CURRENT month's bus charges. Safe to run
# alongside the manual "Generate Monthly Bus Fees" button, or after a crash/
# restart re-runs it - _bus_generation_run() is idempotent per
# (student, academic_year, month, charge_type), so any student already
# charged for the month is simply reported under "already_generated" and
# never charged twice, no matter how many times or how this gets triggered.
async def monthly_bus_charge_scheduler():
    import logging
    log = logging.getLogger("bus_charges.scheduler")
    while True:
        try:
            month = now_iso()[:7]  # 'YYYY-MM'
            result = await _bus_generation_run("2026-27", month, commit=True, user=None)
            if result["generated"]:
                log.info(f"bus charges: auto-generated {result['generated']} charge(s) for {month}")
        except Exception as e:
            log.error(f"bus charges: auto-generation failed: {e}")
        await asyncio.sleep(24 * 60 * 60)  # once a day is enough - generation itself is idempotent
